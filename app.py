import json
import random
import sqlite3
import datetime
import requests
import numpy as np
import streamlit as st
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity
import torch

# =====================================================================
# 1. SQLITE DATABASE MANAGER FOR PERSISTENT STORAGE
# =====================================================================

class DatabaseManager:
    def __init__(self, db_path: str = "learning_assistant.db"):
        self.db_path = db_path
        self._init_db()

    def _init_db(self):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        # Chat History Table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS chat_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                query TEXT,
                answer TEXT,
                timestamp TEXT
            )
        """)
        # Quiz Attempts Table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS quiz_attempts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                question TEXT,
                user_ans TEXT,
                correct_ans TEXT,
                is_correct INTEGER,
                explanation TEXT,
                timestamp TEXT
            )
        """)
        # Tutor Sessions Table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS tutor_sessions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                query TEXT,
                language TEXT,
                explanation TEXT,
                timestamp TEXT
            )
        """)
        # Generated Content Table (Summaries, Question Papers)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS generated_content (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                content_type TEXT,
                content_json TEXT,
                timestamp TEXT
            )
        """)
        conn.commit()
        conn.close()

    def save_chat(self, query: str, answer: str):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute("INSERT INTO chat_history (query, answer, timestamp) VALUES (?, ?, ?)",
                       (query, answer, str(datetime.datetime.now())))
        conn.commit()
        conn.close()

    def save_quiz_attempt(self, q: str, user_ans: str, correct_ans: str, is_correct: bool, explanation: str):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO quiz_attempts (question, user_ans, correct_ans, is_correct, explanation, timestamp)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (q, user_ans, correct_ans, int(is_correct), explanation, str(datetime.datetime.now())))
        conn.commit()
        conn.close()

    def save_tutor_session(self, query: str, language: str, explanation: str):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO tutor_sessions (query, language, explanation, timestamp)
            VALUES (?, ?, ?, ?)
        """, (query, language, explanation, str(datetime.datetime.now())))
        conn.commit()
        conn.close()

    def save_content(self, content_type: str, content_dict: dict):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO generated_content (content_type, content_json, timestamp)
            VALUES (?, ?, ?)
        """, (content_type, json.dumps(content_dict), str(datetime.datetime.now())))
        conn.commit()
        conn.close()

    def load_all_history(self):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        
        # Load Chats
        cursor.execute("SELECT query, answer FROM chat_history ORDER BY id ASC")
        chats = [{"query": r[0], "answer": r[1]} for r in cursor.fetchall()]

        # Load Quiz History & Stats
        cursor.execute("SELECT question, user_ans, correct_ans, is_correct, explanation FROM quiz_attempts ORDER BY id ASC")
        quiz_history = []
        correct_count = 0
        incorrect_count = 0
        for r in cursor.fetchall():
            is_corr = bool(r[3])
            if is_corr:
                correct_count += 1
            else:
                incorrect_count += 1
            quiz_history.append({
                "question": r[0],
                "user_ans": r[1],
                "correct_ans": r[2],
                "is_correct": is_corr,
                "explanation": r[4]
            })

        # Load Tutor Sessions
        cursor.execute("SELECT query, language, explanation FROM tutor_sessions ORDER BY id ASC")
        tutor = [{"query": r[0], "language": r[1], "explanation": r[2]} for r in cursor.fetchall()]

        # Load Latest Summary / Question Paper if available
        cursor.execute("SELECT content_json FROM generated_content WHERE content_type = 'summary' ORDER BY id DESC LIMIT 1")
        sum_row = cursor.fetchone()
        latest_summary = json.loads(sum_row[0]) if sum_row else None

        cursor.execute("SELECT content_json FROM generated_content WHERE content_type = 'question_paper' ORDER BY id DESC LIMIT 1")
        qp_row = cursor.fetchone()
        latest_qp = json.loads(qp_row[0]) if qp_row else None

        conn.close()
        return chats, quiz_history, correct_count, incorrect_count, tutor, latest_summary, latest_qp

    def clear_database(self):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM chat_history")
        cursor.execute("DELETE FROM quiz_attempts")
        cursor.execute("DELETE FROM tutor_sessions")
        cursor.execute("DELETE FROM generated_content")
        conn.commit()
        conn.close()


# =====================================================================
# 2. TEXT EXTRACTION & CHUNKING
# =====================================================================

def extract_pdf_text(uploaded_file) -> str:
    pdf_reader = PdfReader(uploaded_file)
    extracted_text = ""
    for page_num, page in enumerate(pdf_reader.pages):
        page_text = page.extract_text()
        if page_text:
            extracted_text += f"\n[Page {page_num + 1}]\n" + page_text
    return extracted_text.strip()


def extract_docx_text(uploaded_file) -> str:
    doc = Document(uploaded_file)
    parts = []
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            parts.append(f"[Paragraph {idx}] {text}")
    for table_idx, table in enumerate(doc.tables, start=1):
        parts.append(f"[Table {table_idx}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_pptx_text(uploaded_file) -> str:
    prs = Presentation(uploaded_file)
    parts = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        parts.append(f"\n[Slide {slide_num}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
                    parts.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = [shape.text.strip() for shape in slide.notes_slide.notes_text_frame.paragraphs if shape.text.strip()]
            if notes:
                parts.append("[Speaker Notes]")
                parts.extend(notes)
    return "\n".join(parts).strip()


def extract_document_text(uploaded_file, file_type: str) -> str:
    if file_type == "pdf":
        return extract_pdf_text(uploaded_file)
    if file_type == "docx":
        return extract_docx_text(uploaded_file)
    if file_type == "pptx":
        return extract_pptx_text(uploaded_file)
    raise ValueError(f"Unsupported file type: {file_type}")


def chunk_text(text: str, chunk_size: int = 400, overlap: int = 50) -> list:
    words = text.split()
    if not words:
        return []
    chunks = []
    for i in range(0, len(words), chunk_size - overlap):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
    return chunks


# =====================================================================
# 3. LOCAL RAG & OLLAMA LLM INTEGRATION
# =====================================================================

class IntelligentPDFEngine:
    def __init__(self, ollama_model: str = "llama3.2:1b"):
        self.ollama_model = ollama_model
        self.ollama_url = "http://localhost:11434/api/generate"
        
        device = "cuda" if torch.cuda.is_available() else "cpu"
        self.embedder = SentenceTransformer("all-MiniLM-L6-v2", device=device)

    def _call_ollama(self, prompt: str, json_format: bool = False, options: dict = None) -> str:
        payload = {
            "model": self.ollama_model,
            "prompt": prompt,
            "stream": False,
            "options": options or {"num_ctx": 4096}
        }
        if json_format:
            payload["format"] = "json"

        try:
            response = requests.post(self.ollama_url, json=payload, timeout=90)
            if response.status_code == 200:
                return response.json().get("response", "")
            else:
                return "Error: Unable to connect to local Ollama server."
        except Exception as e:
            return f"Error connecting to Ollama: {str(e)}"

    def get_relevant_chunks(self, query: str, chunks: list, top_k: int = 3) -> list:
        if not chunks:
            return []
        query_vec = self.embedder.encode([query])
        chunk_vecs = self.embedder.encode(chunks)
        sims = cosine_similarity(query_vec, chunk_vecs)[0]
        top_indices = np.argsort(sims)[::-1][:top_k]
        return [chunks[i] for i in top_indices]

    def generate_summary(self, full_text: str, progress_callback=None) -> dict:
        words = full_text.split()
        if not words:
            return {"overview": "No readable content found.", "sections": [], "key_takeaways": []}

        words_per_block = 2000
        blocks = [" ".join(words[i:i + words_per_block]) for i in range(0, len(words), words_per_block)]
        total_blocks = len(blocks)
        block_summaries = []

        if total_blocks == 1:
            if progress_callback:
                progress_callback(0.5, "Generating complete document summary...")
            prompt = f"""You are an academic tutor. Create a comprehensive topic-by-topic study guide for this document.

DOCUMENT TEXT:
{blocks[0]}

Return ONLY valid JSON with this exact structure:
{{
  "overview": "Detailed overall overview of the document (4-6 sentences).",
  "sections": [
    {{
      "title": "Topic Name",
      "summary": "Detailed explanation of this topic.",
      "key_points": ["Point 1", "Point 2", "Point 3"]
    }}
  ],
  "key_takeaways": ["Takeaway 1", "Takeaway 2", "Takeaway 3"]
}}"""
            res = self._call_ollama(prompt, json_format=True)
            try:
                return json.loads(res)
            except Exception:
                return {"overview": res, "sections": [], "key_takeaways": []}

        for idx, block in enumerate(blocks, start=1):
            if progress_callback:
                progress_callback((idx / (total_blocks + 1)), f"Processing section block {idx} of {total_blocks}...")

            prompt = f"""Summarize this portion of the document. Cover all key concepts, definitions, and facts.

DOCUMENT BLOCK ({idx}/{total_blocks}):
{block}

SUMMARY OF BLOCK {idx}:"""
            res = self._call_ollama(prompt, json_format=False)
            if res and not res.startswith("Error:"):
                block_summaries.append(f"Section Block {idx}:\n{res.strip()}")

        if progress_callback:
            progress_callback(0.9, "Consolidating final structured summary...")

        combined_blocks = "\n\n====================\n\n".join(block_summaries)
        final_prompt = f"""You are an academic tutor creating a structured study guide from the following section summaries.

SECTION SUMMARIES:
{combined_blocks}

Return ONLY valid JSON matching this exact structure:
{{
  "overview": "Comprehensive 4-6 sentence overview of the whole document.",
  "sections": [
    {{
      "title": "Topic Name",
      "summary": "Detailed explanation of this topic.",
      "key_points": ["Point 1", "Point 2", "Point 3"]
    }}
  ],
  "key_takeaways": ["Takeaway 1", "Takeaway 2", "Takeaway 3"]
}}"""

        res = self._call_ollama(final_prompt, json_format=True)
        try:
            return json.loads(res)
        except Exception:
            return {
                "overview": "Section-by-section breakdown compiled below:",
                "sections": [{"title": f"Block {idx+1}", "summary": text, "key_points": []} for idx, text in enumerate(block_summaries)],
                "key_takeaways": []
            }

    def generate_quiz(self, chunks: list, num_questions: int = 4) -> list:
        if not chunks:
            return []
        valid_chunks = [c for c in chunks if len(c.split()) >= 60] or chunks
        sampled_chunks = random.sample(valid_chunks, min(num_questions, len(valid_chunks)))
        validated_questions = []

        for passage in sampled_chunks:
            prompt = f"""You are a university professor creating a unique exam question based strictly on the passage below.

READ THIS PASSAGE CAREFULLY:
{passage}

INSTRUCTIONS:
1. Create ONE clear, realistic multiple-choice question testing a main concept explained in the passage.
2. Provide EXACTLY 4 options (1 correct, 3 plausible distractors).

Return ONLY valid JSON matching this exact format:
{{
  "question": "What is ...?",
  "options": ["Option A", "Option B", "Option C", "Option D"],
  "correct_index": 0,
  "explanation": "State clearly why the correct option is right based on the text."
}}"""
            response_text = self._call_ollama(prompt, json_format=True, options={"num_ctx": 4096, "temperature": 0.8})
            clean_text = response_text.strip()
            if "```" in clean_text:
                parts = clean_text.split("```")
                for part in parts:
                    if part.strip().startswith("json"):
                        clean_text = part.strip()[4:].strip()
                        break
                    elif part.strip().startswith("{"):
                        clean_text = part.strip()
                        break
            try:
                data = json.loads(clean_text)
                q_text = data.get("question", "").strip()
                opts = [str(o).strip() for o in data.get("options", []) if str(o).strip()]
                c_idx = int(data.get("correct_index", 0))
                if q_text and len(opts) >= 4:
                    validated_questions.append({
                        "question": q_text,
                        "options": opts[:4],
                        "correct_index": max(0, min(c_idx, 3)),
                        "explanation": str(data.get("explanation", "Grounded in source passage.")).strip()
                    })
            except Exception:
                continue
        return validated_questions

    def generate_question_paper(self, chunks: list, pyq_text: str = "") -> dict:
        if not chunks:
            return {}
        valid_chunks = [c for c in chunks if len(c.split()) >= 50] or chunks
        sampled_chunks = random.sample(valid_chunks, min(6, len(valid_chunks)))
        selected_text = "\n\n---\n\n".join(sampled_chunks)

        pyq_instruction = f"\nPREVIOUS YEARS' PAPERS:\n{pyq_text[:4000]}\nPrioritize recurring themes." if pyq_text.strip() else ""

        prompt = f"""You are an expert academic examiner framing a formal examination question paper.
{pyq_instruction}
SOURCE MATERIAL:
{selected_text}
REQUIREMENTS:
- Section A: 5 short questions worth 1 Mark each.
- Section B: 3 conceptual questions worth 3 Marks each.
- Section C: 3 descriptive questions worth 5 Marks each.

Return ONLY valid JSON matching this schema:
{{
  "title": "Examination Question Paper",
  "section_a": ["Q1?", "Q2?", "Q3?", "Q4?", "Q5?"],
  "section_b": ["Q6?", "Q7?", "Q8?"],
  "section_c": ["Q9?", "Q10?", "Q11?"]
}}"""
        response_text = self._call_ollama(prompt, json_format=True, options={"num_ctx": 4096, "temperature": 0.7})
        clean_text = response_text.strip()
        if "```" in clean_text:
            parts = clean_text.split("```")
            for part in parts:
                if part.strip().startswith("json"):
                    clean_text = part.strip()[4:].strip()
                    break
                elif part.strip().startswith("{"):
                    clean_text = part.strip()
                    break
        try:
            data = json.loads(clean_text)
            return {
                "title": str(data.get("title", "Examination Question Paper")).strip(),
                "section_a": [str(q).strip() for q in data.get("section_a", []) if str(q).strip()][:5],
                "section_b": [str(q).strip() for q in data.get("section_b", []) if str(q).strip()][:3],
                "section_c": [str(q).strip() for q in data.get("section_c", []) if str(q).strip()][:3]
            }
        except Exception:
            return {}

    def answer_question(self, user_question: str, context_chunks: list) -> str:
        context_str = "\n\n---\n\n".join(context_chunks)
        prompt = f"""Answer the question using ONLY the provided document context. If not found, say "I couldn't find that in the document."

Context:
{context_str}

Question: {user_question}
Answer:"""
        return self._call_ollama(prompt, json_format=False)

    def tutor_explain(self, query: str, context_chunks: list, language: str = "Hindi") -> str:
        context_str = "\n\n---\n\n".join(context_chunks)
        prompt = f"""You are a helpful AI tutor helping a student understand their study materials. 
Explain the concept in fluent {language} using a conversational teaching tone based strictly on the context.

Context Passages:
{context_str}

Student Question / Concept: {query}
Tutor Explanation ({language}):"""
        return self._call_ollama(prompt, json_format=False, options={"num_ctx": 4096, "temperature": 0.7})


# =====================================================================
# 4. STREAMLIT INTERFACE & DATABASE SYNCHRONIZATION
# =====================================================================

def main():
    st.set_page_config(page_title="EduAI Studio — On-Device Learning Assistant", page_icon="🎓", layout="wide")

    # Initialize Database Manager
    if "db" not in st.session_state:
        st.session_state.db = DatabaseManager()

    # Initialize Session States from Database
    if "initialized_from_db" not in st.session_state:
        chats, quiz_history, correct_count, incorrect_count, tutor, saved_summary, saved_qp = st.session_state.db.load_all_history()
        
        st.session_state.chat_history = chats
        st.session_state.tutor_history = tutor
        st.session_state.quiz_stats = {
            "total_answered": correct_count + incorrect_count,
            "correct": correct_count,
            "incorrect": incorrect_count,
            "history": quiz_history
        }
        st.session_state.answered_questions = {f"{item['question']}" for item in quiz_history}
        
        if saved_summary:
            st.session_state.summary_data = saved_summary
        if saved_qp:
            st.session_state.qp_data = saved_qp

        st.session_state.initialized_from_db = True

    if "engine" not in st.session_state:
        st.session_state.engine = IntelligentPDFEngine()
    if "pdf_text" not in st.session_state:
        st.session_state.pdf_text = None
    if "pdf_chunks" not in st.session_state:
        st.session_state.pdf_chunks = []
    if "file_names" not in st.session_state:
        st.session_state.file_names = []

    # --- SIDEBAR UI ---
    st.sidebar.markdown("## 🎓 EduAI Studio")
    st.sidebar.caption("Privacy-First Local Learning Assistant with SQLite Persistence")
    st.sidebar.markdown("---")

    with st.sidebar.expander("⚙️ Model Configuration", expanded=False):
        model_name = st.text_input("Ollama Model Name", value="llama3.2:1b")
        if model_name != st.session_state.engine.ollama_model:
            st.session_state.engine.ollama_model = model_name

    st.sidebar.markdown("### 📁 Document Library")
    uploaded_files = st.sidebar.file_uploader(
        "Upload Study Files", 
        type=["pdf", "docx", "pptx"], 
        accept_multiple_files=True
    )

    if uploaded_files:
        if st.sidebar.button("✨ Process Files & Index", type="primary", use_container_width=True):
            with st.spinner("Extracting & embedding documents..."):
                combined_text = ""
                file_names = []
                for uploaded_file in uploaded_files:
                    file_type = uploaded_file.name.rsplit(".", 1)[-1].lower()
                    text = extract_document_text(uploaded_file, file_type)
                    if text:
                        combined_text += f"\n\n=== FILE: {uploaded_file.name} ===\n\n" + text
                        file_names.append(uploaded_file.name)

                if not combined_text.strip():
                    st.sidebar.error("No readable text found in uploaded files.")
                else:
                    chunks = chunk_text(combined_text)
                    st.session_state.pdf_text = combined_text
                    st.session_state.pdf_chunks = chunks
                    st.session_state.file_names = file_names
                    st.sidebar.success(f"Successfully indexed {len(file_names)} file(s)!")
                    st.toast("Documents processed successfully!", icon="🚀")

    if st.sidebar.button("🗑️ Clear Database History", use_container_width=True):
        st.session_state.db.clear_database()
        st.session_state.chat_history = []
        st.session_state.tutor_history = []
        st.session_state.quiz_stats = {"total_answered": 0, "correct": 0, "incorrect": 0, "history": []}
        st.session_state.answered_questions = set()
        st.session_state.pop("summary_data", None)
        st.session_state.pop("qp_data", None)
        st.toast("Database history cleared!", icon="🗑️")
        st.rerun()

    # --- MAIN VIEW HEADER ---
    st.markdown("# 🧠 AI-Powered On-Device Learning Studio")
    
    if not st.session_state.pdf_text:
        st.info("👋 **Welcome!** Please upload one or more study documents (`PDF`, `DOCX`, or `PPTX`) from the left sidebar and click **Process Files & Index** to begin.")
        return

    st.success(f"📂 **Active Library:** {', '.join(st.session_state.file_names)} — *{len(st.session_state.pdf_text.split()):,} total words indexed*")

    # --- TABS ---
    tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
        "📖 Summary", 
        "❓ Practice Quiz", 
        "📝 Question Paper", 
        "💬 Document Q&A",
        "🗣️ AI Tutor",
        "📊 Dashboard"
    ])

    # --- TAB 1: SUMMARY ---
    with tab1:
        st.markdown("### 📖 Comprehensive Topic-by-Topic Summary")
        if st.button("✨ Generate Full Document Summary", type="primary"):
            p_bar = st.progress(0.0, text="Analyzing documents...")
            def update_progress(val, msg):
                p_bar.progress(val, text=msg)

            summary = st.session_state.engine.generate_summary(st.session_state.pdf_text, progress_callback=update_progress)
            st.session_state.summary_data = summary
            st.session_state.db.save_content("summary", summary)
            p_bar.empty()
            st.toast("Summary saved to database!", icon="💾")

        if "summary_data" in st.session_state:
            s = st.session_state.summary_data
            with st.container():
                st.markdown("#### 📌 Executive Overview")
                st.write(s.get("overview", "No overview generated."))

            sections = s.get("sections", [])
            if sections:
                st.markdown("#### 📚 Detailed Topic Breakdown")
                for idx, section in enumerate(sections, start=1):
                    with st.container():
                        st.markdown(f"**{idx}. {section.get('title', 'Topic')}**")
                        st.write(section.get("summary", ""))
                        points = section.get("key_points", [])
                        if points:
                            st.markdown("*Key Takeaways:*")
                            for point in points:
                                st.markdown(f"- {point}")

            with st.container():
                st.markdown("#### 🎯 Core Study Takeaways")
                for pt in s.get("key_takeaways", []):
                    st.markdown(f"- {pt}")

    # --- TAB 2: PRACTICE QUIZ ---
    with tab2:
        st.markdown("### ❓ Adaptive Practice Quiz")
        if st.button("🎲 Generate Fresh Quiz Questions", type="primary"):
            with st.spinner("Sampling library and building questions..."):
                questions = st.session_state.engine.generate_quiz(st.session_state.pdf_chunks, num_questions=4)
                st.session_state.quiz_data = questions
            st.toast("New quiz generated!", icon="🎲")

        if "quiz_data" in st.session_state:
            questions = st.session_state.quiz_data
            for idx, q in enumerate(questions):
                with st.container():
                    st.markdown(f"**Question {idx + 1}:** {q['question']}")
                    options = q.get("options", [])
                    if options:
                        user_ans = st.radio(f"Select answer for Q{idx + 1}:", options, key=f"q_radio_{idx}")
                        if st.button(f"Submit Answer #{idx + 1}", key=f"q_submit_{idx}"):
                            selected_idx = options.index(user_ans)
                            correct_idx = q.get("correct_index", 0)
                            is_correct = (selected_idx == correct_idx)
                            
                            q_key = f"{q['question']}_{idx}"
                            if q_key not in st.session_state.answered_questions:
                                st.session_state.answered_questions.add(q_key)
                                st.session_state.quiz_stats["total_answered"] += 1
                                if is_correct:
                                    st.session_state.quiz_stats["correct"] += 1
                                else:
                                    st.session_state.quiz_stats["incorrect"] += 1

                                st.session_state.quiz_stats["history"].append({
                                    "question": q["question"],
                                    "user_ans": user_ans,
                                    "correct_ans": options[correct_idx],
                                    "is_correct": is_correct,
                                    "explanation": q.get("explanation", "")
                                })
                                # Save attempt to SQLite database
                                st.session_state.db.save_quiz_attempt(
                                    q["question"], user_ans, options[correct_idx], is_correct, q.get("explanation", "")
                                )
                                st.toast("Quiz attempt saved to database!", icon="💾")

                            if is_correct:
                                st.success("🎉 Spot on! Correct Answer.")
                            else:
                                st.error(f"❌ Incorrect. Correct answer was: **{options[correct_idx]}**")
                            st.info(f"**Explanation:** {q.get('explanation', '')}")

    # --- TAB 3: QUESTION PAPER & PYQ ANALYSIS ---
    with tab3:
        st.markdown("### 📝 Formal Examination Question Paper Generator")
        pyq_files = st.file_uploader("Upload PYQs (Optional)", type=["pdf", "docx", "pptx"], accept_multiple_files=True, key="pyq_files_uploader")
        
        pyq_text_content = ""
        if pyq_files:
            for pyq in pyq_files:
                f_type = pyq.name.rsplit(".", 1)[-1].lower()
                t_content = extract_document_text(pyq, f_type)
                if t_content:
                    pyq_text_content += f"\n\n--- PYQ FILE: {pyq.name} ---\n\n" + t_content

        if st.button("📜 Generate Exam Paper (5x1, 3x3, 3x5)", type="primary"):
            with st.spinner("Analyzing themes and framing question paper..."):
                qp = st.session_state.engine.generate_question_paper(st.session_state.pdf_chunks, pyq_text=pyq_text_content)
                st.session_state.qp_data = qp
                st.session_state.db.save_content("question_paper", qp)
            st.toast("Question paper saved to database!", icon="💾")

        if "qp_data" in st.session_state:
            qp = st.session_state.qp_data
            if qp and qp.get("section_a"):
                sec_a = qp.get("section_a", [])
                sec_b = qp.get("section_b", [])
                sec_c = qp.get("section_c", [])

                with st.container():
                    st.markdown(f"<h3 style='text-align: center;'>{qp.get('title', 'EXAMINATION QUESTION PAPER')}</h3>", unsafe_allow_html=True)
                    col1, col2, col3 = st.columns(3)
                    col1.metric("Time Allowed", "1 Hour")
                    col2.metric("Total Questions", "11 Questions")
                    col3.metric("Max Marks", "29 Marks")
                    st.markdown("---")

                    st.markdown("#### SECTION A: Short Answer Questions (5 x 1 = 5 Marks)")
                    for idx, q in enumerate(sec_a, start=1): st.markdown(f"**Q{idx}.** {q} `[1 Mark]`")
                    st.markdown("---")
                    st.markdown("#### SECTION B: Medium Conceptual Questions (3 x 3 = 9 Marks)")
                    for idx, q in enumerate(sec_b, start=6): st.markdown(f"**Q{idx}.** {q} `[3 Marks]`")
                    st.markdown("---")
                    st.markdown("#### SECTION C: Descriptive / Essay Questions (3 x 5 = 15 Marks)")
                    for idx, q in enumerate(sec_c, start=9): st.markdown(f"**Q{idx}.** {q} `[5 Marks]`")

                paper_text = f"{qp.get('title', 'EXAMINATION QUESTION PAPER')}\nTime Allowed: 1 Hour | Total Marks: 29 Marks\n" + "="*50 + "\n\n"
                paper_text += "SECTION A (5 x 1 = 5 Marks)\n" + "\n".join([f"Q{i+1}. {q} [1M]" for i, q in enumerate(sec_a)]) + "\n\n"
                paper_text += "SECTION B (3 x 3 = 9 Marks)\n" + "\n".join([f"Q{i+6}. {q} [3M]" for i, q in enumerate(sec_b)]) + "\n\n"
                paper_text += "SECTION C (3 x 5 = 15 Marks)\n" + "\n".join([f"Q{i+9}. {q} [5M]" for i, q in enumerate(sec_c)])

                st.download_button(
                    label="📥 Download Question Paper (.txt)",
                    data=paper_text,
                    file_name="Exam_Question_Paper.txt",
                    mime="text/plain"
                )

    # --- TAB 4: DOCUMENT Q&A ---
    with tab4:
        st.markdown("### 💬 Context-Grounded Document Q&A")
        user_query = st.text_input("Ask a question:", placeholder="e.g., What is the main conclusion?")
        
        if st.button("🔍 Search & Answer", type="primary"):
            if user_query:
                with st.spinner("Searching document library..."):
                    context_chunks = st.session_state.engine.get_relevant_chunks(user_query, st.session_state.pdf_chunks)
                    answer = st.session_state.engine.answer_question(user_query, context_chunks)
                    st.session_state.chat_history.append({"query": user_query, "answer": answer, "sources": context_chunks})
                    # Save chat to SQLite database
                    st.session_state.db.save_chat(user_query, answer)
                    st.toast("Chat saved to database!", icon="💾")

        if st.session_state.chat_history:
            st.markdown("---")
            for item in reversed(st.session_state.chat_history):
                with st.container():
                    st.markdown(f"**Q:** {item['query']}")
                    st.markdown(f"**Answer:**\n{item['answer']}")

    # --- TAB 5: AI TUTOR (MULTILINGUAL) ---
    with tab5:
        st.markdown("### 🗣️ Multilingual AI Tutor")
        t_col1, t_col2 = st.columns([2, 1])
        with t_col2:
            target_language = st.selectbox("Select Language / भाषा चुनें:", ["Hindi (हिंदी)", "Hinglish", "English", "Spanish", "French", "German"])
        with t_col1:
            tutor_query = st.text_input("What would you like explained?", placeholder="e.g., Explain the core mechanism...")

        if st.button("🎓 Explain Concept", type="primary"):
            if tutor_query:
                with st.spinner(f"Preparing tutor explanation in {target_language}..."):
                    context_chunks = st.session_state.engine.get_relevant_chunks(tutor_query, st.session_state.pdf_chunks, top_k=3)
                    explanation = st.session_state.engine.tutor_explain(tutor_query, context_chunks, language=target_language)
                    st.session_state.tutor_history.append({"query": tutor_query, "language": target_language, "explanation": explanation, "sources": context_chunks})
                    # Save tutor session to SQLite database
                    st.session_state.db.save_tutor_session(tutor_query, target_language, explanation)
                    st.toast("Tutor session saved to database!", icon="💾")

        if st.session_state.tutor_history:
            st.markdown("---")
            for item in reversed(st.session_state.tutor_history):
                with st.container():
                    st.markdown(f"**Topic:** {item['query']} *({item['language']})*")
                    st.markdown(item["explanation"])

    # --- TAB 6: DASHBOARD ---
    with tab6:
        st.markdown("### 📊 Quiz Performance Dashboard")
        stats = st.session_state.quiz_stats
        total = stats["total_answered"]
        correct = stats["correct"]
        incorrect = stats["incorrect"]
        accuracy = (correct / total * 100) if total > 0 else 0.0

        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Total Answered", total)
        m2.metric("Correct", correct)
        m3.metric("Incorrect", incorrect)
        m4.metric("Accuracy Rate", f"{accuracy:.1f}%")

        st.markdown("---")
        st.markdown("#### 🎯 Overall Mastery Rank")
        st.progress(accuracy / 100.0, text=f"Mastery Level: {accuracy:.1f}%")

        st.markdown("---")
        st.markdown("#### 📋 Attempt History Log (Persistent)")
        if stats["history"]:
            for idx, item in enumerate(reversed(stats["history"]), start=1):
                status_label = "✅ Correct" if item["is_correct"] else "❌ Incorrect"
                with st.expander(f"Attempt #{idx}: {item['question'][:50]}... ({status_label})"):
                    st.markdown(f"**Question:** {item['question']}")
                    st.markdown(f"**Your Answer:** {item['user_ans']}")
                    st.markdown(f"**Correct Answer:** {item['correct_ans']}")
                    st.markdown(f"**Explanation:** {item['explanation']}")
        else:
            st.caption("No quiz attempts recorded yet.")

if __name__ == "__main__":
    main()  